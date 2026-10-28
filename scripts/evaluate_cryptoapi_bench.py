#!/usr/bin/env python3
"""Print location-based CryptoAPI-Bench metrics; never modify reports or ground truth.

TP: vulnerable case with an annotated line inside a reported range.
FN: vulnerable case with no findings.
FP + FN: vulnerable case with findings, but no matching location.
FP: non-vulnerable case with findings. TN: non-vulnerable case without findings.
One spreadsheet row is one case per analysis directory, not per finding.
"""
    
import argparse
import datetime
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


# Read Excel with openpyxl and recover line lists that Excel converted into dates.
def read_spreadsheet(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Vulnerabilities"]
        headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))]
        line_column = headers.index("Line number")
        rows = []
        for number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            record = {header: "" if cell.value is None else str(cell.value)
                      for header, cell in zip(headers, cells) if header}
            if not record.get("Files", "").strip():
                continue

            record.update(row=number, date_recovered=False)
            cell = cells[line_column]
            number_format = cell.number_format.replace('\\', "").replace(" ", "")
            if isinstance(cell.value, datetime.date) and number_format in ("m,d", "m,d,yy"):
                lines = [cell.value.month, cell.value.day]
                if number_format.endswith(",yy"):
                    lines.append(cell.value.year % 100)

                record["Line number"] = ",".join(map(str, lines))
                record["date_recovered"] = True

            rows.append(record)

        return rows
    finally:
        workbook.close()


# Keep comma-separated lines distinct; parentheses group lines for one file.
def line_groups(text):
    if not text.strip() or text.strip() == "---":
        return []

    if not re.fullmatch(r"[\d\s,();\-–]+", text):
        raise ValueError(f"Unrecognized line notation: {text!r}")

    groups = []
    for token in re.findall(r"\([^()]*\)|\d+(?:\s*[-–]\s*\d+)?", text):
        numbers = []
        for part in token.strip("()").split(","):
            ends = re.split(r"\s*[-–]\s*", part.strip())
            start, end = int(ends[0]), int(ends[-1])
            if start < 1 or end < start:
                raise ValueError(f"Invalid line numbers: {token!r}")

            numbers.extend(range(start, end + 1))

        groups.append(numbers)

    return groups


# Find the actual analysis directory of each Java file, including bundled inputs.
def analysis_directories(run):
    locations = defaultdict(list)
    for source in sorted(run.glob("*/source")):
        for path in source.rglob("*.java"):
            if "security-scan" not in path.relative_to(source).parts:
                locations[path.name].append(source.parent)

    return locations


# Use the XLSX label and inclusive reported range; no tolerance or type matching.
def classify(vulnerable, has_findings, matched):
    if not vulnerable:
        return ["FP"] if has_findings else ["TN"]

    if matched:
        return ["TP"]

    return ["FP", "FN"] if has_findings else ["FN"]


# Group each spreadsheet row by where its files were actually scanned.
def evaluate(rows, locations):
    results = []
    issues = []
    annotated_files = set()
    report_cache = {}
    for row in rows:
        names = [name.strip() for name in row["Files"].splitlines() if name.strip()]
        names = [name if name.endswith(".java") else name + ".java" for name in names]
        annotated_files.update(names)
        flag = str(row.get("Vulnerability Exists?", "")).strip().lower()
        if flag not in ("", "0", "1", "false", "true"):
            issues.append({"row": row["row"], "files": names, "issue": "invalid vulnerability flag"})
            continue

        vulnerable = flag in ("1", "true")
        if not flag:
            issues.append({"row": row["row"], "files": names,
                           "issue": "blank vulnerability flag treated as false, per user instruction",
                           "line_cell": row.get("Line number", "")})

        try:
            groups = line_groups(row.get("Line number", "")) if vulnerable else []
        except ValueError as error:
            issues.append({"row": row["row"], "files": names, "issue": str(error)})
            continue

        if len(names) == 1:
            assignments = [sorted({number for group in groups for number in group})]
        elif len(groups) <= len(names):
            assignments = groups + [None] * (len(names) - len(groups))
        else:
            issues.append({"row": row["row"], "files": names, "issue": "ambiguous line-to-file mapping"})
            continue

        by_directory = defaultdict(dict)
        for index, name in enumerate(names):
            directories = locations.get(name, [])
            if len(directories) != 1:
                issues.append({"row": row["row"], "files": [name],
                               "issue": "no analysis found" if not directories else "multiple analyses found"})
                continue

            by_directory[directories[0]][name] = assignments[index] if vulnerable else []

        for directory, expected in by_directory.items():
            if vulnerable and not any(expected.values()):
                issues.append({"row": row["row"], "files": list(expected),
                               "issue": "no unambiguous ground-truth lines for this analysis"})
                continue

            report_path = directory / "source/security-scan/findings.json"
            try:
                if report_path not in report_cache:
                    report = json.loads(report_path.read_text(encoding="utf-8"))
                    if not isinstance(report.get("findings", []), list):
                        raise ValueError("findings must be a list")

                    report_cache[report_path] = report

                report = report_cache[report_path]
            except (OSError, ValueError) as error:
                issues.append({"row": row["row"], "files": list(expected),
                               "issue": f"report unavailable: {error}"})
                continue

            predictions = []
            for ranked in report.get("findings", []):
                finding = ranked.get("finding", ranked)
                name = PurePosixPath(str(finding.get("file", "")).replace("\\", "/")).name
                if name in expected:
                    predictions.append(finding)

            matched = []
            ranges = []
            for finding in predictions:
                name = PurePosixPath(finding["file"].replace("\\", "/")).name
                start = finding.get("line_start")
                end = finding.get("line_end")
                ranges.append({"file": name, "start": start, "end": end})
                if type(start) is not int or type(end) is not int or start < 1 or end < start:
                    issues.append({"row": row["row"], "files": [name], "issue": "invalid predicted line range"})
                    continue

                matched.extend({"file": name, "line": number} for number in expected[name] or []
                               if start <= number <= end)

            status_path = directory / "status.json"
            status = json.loads(status_path.read_text()) if status_path.exists() else {}
            error_logs = list((directory / "source/security-scan").glob("*errors*.jsonl"))
            results.append({
                "spreadsheet_row": row["row"], "analysis": directory.name,
                "files": list(expected), "vulnerable": vulnerable,
                "ground_truth_lines": expected, "reported_ranges": ranges, "matched_lines": matched,
                "counts": classify(vulnerable, bool(predictions), bool(matched)),
                "status": status.get("status", "unknown"), "degraded": report.get("degraded"),
                "error_log_present": any(path.stat().st_size > 0 for path in error_logs),
                "report": str(report_path),
            })

    return results, issues, annotated_files


# Print metrics and optional per-case details; create no result or cache files.
def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-dir", type=Path, help="Dataset session directory; default: latest saved session")
    parser.add_argument("--xlsx", type=Path, default=ROOT / "datasets/CryptoAPI-Bench/CryptoAPI-Bench_details.xlsx")
    parser.add_argument("--details", action="store_true", help="Include each case's lines, predictions, and counts")
    args = parser.parse_args()
    sessions = sorted(path for path in (ROOT / "runs/cryptoapi-bench").glob("20*-*") if path.is_dir())
    if args.run_dir is None and not sessions:
        parser.error("No saved dataset sessions found; pass --run-dir")

    run = (args.run_dir or sessions[-1]).resolve()
    if not run.is_dir():
        parser.error(f"Run directory does not exist: {run}")

    rows = read_spreadsheet(args.xlsx)
    locations = analysis_directories(run)
    results, issues, annotated_files = evaluate(rows, locations)
    counts = Counter({name: 0 for name in ("TP", "FP", "FN", "TN")})
    for case in results:
        counts.update(case["counts"])

    tp, fp, fn = counts["TP"], counts["FP"], counts["FN"]
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else 0.0
    summary = {
        "run_directory": str(run), "ground_truth": str(args.xlsx.resolve()),
        "ground_truth_sha256": hashlib.sha256(args.xlsx.read_bytes()).hexdigest(),
        "counts": dict(counts), "precision": precision, "recall": recall, "f1": f1,
        "percentages": {"precision": round(100 * precision, 2), "recall": round(100 * recall, 2), "f1": round(100 * f1, 2)},
        "coverage": {
            "spreadsheet_rows": len(rows), "analysed_files": len(locations), "scored_cases": len(results),
            "scored_vulnerable_cases": sum(case["vulnerable"] for case in results),
            "scored_non_vulnerable_cases": sum(not case["vulnerable"] for case in results),
            "analysed_files_without_spreadsheet_entry": sorted(set(locations) - annotated_files),
            "analysed_files_not_scored": sorted(set(locations) - {name for case in results for name in case["files"]}),
            "scored_cases_with_error_logs": sum(case["error_log_present"] for case in results),
            "date_cells_recovered": sum(row["date_recovered"] for row in rows),
        },
        "issues": issues,
    }
    if args.details:
        summary["cases"] = results

    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
